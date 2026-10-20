"""Table tools: layout detection, table discovery, full extraction."""

from __future__ import annotations

from typing import Any, Dict, List, Optional, Tuple
import re

from ..core.schemas import ExtractionMethod, SourceType
from ..tables.logical import fragments_from_grid, reconstruct
from ..workspace.world_model import WorldModel
from .registry import ToolError, provenance
from .text import _require_doc


def detect_layout(
    world: WorldModel, doc_id: str, page: Optional[int] = None
) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    _, parsed = _require_doc(world, doc_id)
    pages = [p for p in parsed.pages if page is None or p.page_no == page]
    if page is not None and not pages:
        raise ToolError(f"document '{doc_id}' has no page {page}")
    layout = [
        {
            "page": p.page_no,
            "char_count": p.char_count,
            "blocks": len(p.blocks),
            "headings": len(p.headings),
            "tables": len(p.tables),
            "image_coverage": round(p.image_coverage, 3),
            "has_text": p.char_count >= 30,
        }
        for p in pages
    ]
    prov = [
        provenance(
            doc_id,
            page=p.page_no,
            source_type=SourceType.PAGE,
            method=ExtractionMethod.NATIVE_TEXT,
        )
        for p in pages
    ]
    return ({"doc_id": doc_id, "pages": layout}, prov)


def detect_tables(
    world: WorldModel, doc_id: str, page: Optional[int] = None
) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    _, parsed = _require_doc(world, doc_id)
    found = []
    prov = []
    for p in parsed.pages:
        if page is not None and p.page_no != page:
            continue
        for idx, table in enumerate(p.tables):
            header = next(([c for c in row if c] for row in table.rows if any(row)), [])
            found.append(
                {
                    "page": p.page_no,
                    "index": idx,
                    "rows": table.n_rows,
                    "cols": table.n_cols,
                    "header": header[:8],
                    "source": table.source,
                }
            )
            prov.append(
                provenance(
                    doc_id,
                    page=p.page_no,
                    source_type=SourceType.TABLE,
                    method=ExtractionMethod.NATIVE_TABLE,
                    snippet=" | ".join(header[:8]),
                )
            )
    for sheet in parsed.sheets:
        if sheet.max_row > 1:
            found.append(
                {
                    "page": None,
                    "sheet": sheet.name,
                    "index": len(found),
                    "rows": sheet.max_row,
                    "cols": sheet.max_col,
                    "merged": list(sheet.merged),
                    "source": "spreadsheet",
                }
            )
            prov.append(
                provenance(
                    doc_id,
                    cell=f"{sheet.name}!A1",
                    source_type=SourceType.TABLE,
                    method=ExtractionMethod.SPREADSHEET,
                )
            )
    return ({"doc_id": doc_id, "count": len(found), "tables": found}, prov)


def extract_table(
    world: WorldModel,
    doc_id: str,
    table_index: int = 0,
    page: Optional[int] = None,
    sheet: Optional[str] = None,
) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    _, parsed = _require_doc(world, doc_id)
    if sheet:
        target = next(
            (s for s in parsed.sheets if s.name.lower() == sheet.lower()), None
        )
        if target is None:
            raise ToolError(f"sheet '{sheet}' not found in '{doc_id}'")
        from openpyxl.utils import get_column_letter

        rows = []
        for r in range(1, target.max_row + 1):
            rows.append([target.cell_text(r, c) for c in range(1, target.max_col + 1)])
        prov = [
            provenance(
                doc_id,
                cell=f"{target.name}!A1:{get_column_letter(target.max_col)}{target.max_row}",
                source_type=SourceType.TABLE,
                method=ExtractionMethod.SPREADSHEET,
            )
        ]
        return (
            {
                "doc_id": doc_id,
                "sheet": target.name,
                "rows": rows,
                "n_rows": len(rows),
                "n_cols": target.max_col,
            },
            prov,
        )
    tables = [
        (p, idx, t)
        for p in parsed.pages
        for idx, t in enumerate(p.tables)
        if page is None or p.page_no == page
    ]
    if not tables or not (0 <= table_index < len(tables)):
        raise ToolError(
            f"table {table_index} not found in '{doc_id}'"
            + (f" page {page}" if page else "")
        )
    p, _, table = tables[table_index]
    prov = [
        provenance(
            doc_id,
            page=p.page_no,
            source_type=SourceType.TABLE,
            method=ExtractionMethod.NATIVE_TABLE,
            snippet=" | ".join(table.rows[0]) if table.rows else "",
        )
    ]
    return (
        {
            "doc_id": doc_id,
            "page": p.page_no,
            "index": table_index,
            "rows": table.rows,
            "n_rows": table.n_rows,
            "n_cols": table.n_cols,
        },
        prov,
    )


def _is_number(cell: str) -> bool:
    return bool(re.fullmatch(r"-?\d[\d,]*(?:\.\d+)?%?", (cell or "").strip()))


def _header_depth(rows: List[List[str]]) -> int:
    """Leading mostly-non-numeric rows are headers (cap 3, min 1)."""
    depth = 0
    for row in rows[:3]:
        cells = [c for c in row if str(c).strip()]
        if cells and sum(1 for c in cells if not _is_number(c)) >= len(cells) / 2:
            depth += 1
        else:
            break
    return max(1, depth)


def _sheet_banner(values: List[List[Any]]) -> tuple:
    """Detect a banner/group row above the header row in sheet grids.

    Delegates to tables.logical.detect_header_layout (single source of truth).
    """
    from ..tables.logical import detect_header_layout

    grid = [[str(c) for c in row] for row in values]
    return detect_header_layout(grid)


def reconstruct_table(
    world: WorldModel,
    doc_id: str,
    page: Optional[int] = None,
    sheet: Optional[str] = None,
) -> Tuple[Dict[str, Any], List[Dict[str, Any]]]:
    """Native fragments -> logical tables with header paths + continuations."""
    from openpyxl.utils import get_column_letter

    _, parsed = _require_doc(world, doc_id)
    fragments = []
    prov: List[Dict[str, Any]] = []
    if sheet:
        target = next(
            (s for s in parsed.sheets if s.name.lower() == sheet.lower()), None
        )
        if target is None:
            raise ToolError(f"sheet '{sheet}' not found in '{doc_id}'")
        grid = [
            [target.values.get((r, c), "") for c in range(1, target.max_col + 1)]
            for r in range(1, target.max_row + 1)
        ]
        grid = [[("" if v is None else str(v)) for v in row] for row in grid]
        banner, header_at = _sheet_banner(grid)
        headers = grid[header_at : header_at + 1] if grid else []
        body = grid[header_at + 1 :]
        frag = fragments_from_grid(
            doc_id,
            None,
            target.name,
            0,
            headers + body,
            header_depth=1,
            method=ExtractionMethod.SPREADSHEET,
            banner=banner,
            bbox=None,
        )
        fragments.append(frag)
        prov.append(
            provenance(
                doc_id,
                cell=f"{target.name}!A1:{get_column_letter(target.max_col)}"
                f"{target.max_row}",
                source_type=SourceType.TABLE,
                method=ExtractionMethod.SPREADSHEET,
            )
        )
    else:
        pages = [p for p in parsed.pages if page is None or p.page_no == page]
        if page is not None and not pages:
            raise ToolError(f"document '{doc_id}' has no page {page}")
        for p in pages:
            for idx, table in enumerate(p.tables):
                frag = fragments_from_grid(
                    doc_id,
                    p.page_no,
                    None,
                    idx,
                    table.rows,
                    header_depth=_header_depth(table.rows),
                    method=ExtractionMethod.NATIVE_TABLE,
                )
                fragments.append(frag)
                prov.append(
                    provenance(
                        doc_id,
                        page=p.page_no,
                        source_type=SourceType.TABLE,
                        method=ExtractionMethod.NATIVE_TABLE,
                    )
                )
    if not fragments:
        raise ToolError(f"no native tables to reconstruct in '{doc_id}'")
    logical = reconstruct(fragments)
    data_tables = []
    for table in logical:
        rows_out = []
        for r in table.rows:
            frag = (
                table.fragments[r["fragment"]]
                if r["fragment"] < len(table.fragments)
                else None
            )
            rows_out.append(
                {
                    "values": r["values"],
                    "fragment": r["fragment"],
                    "row": r["row"],
                    "page": frag.page if frag else None,
                    "sheet": frag.sheet if frag else None,
                }
            )
        data_tables.append(
            {
                "title": table.title,
                "header_paths": table.header_paths,
                "rows": rows_out,
                "fragments": len(table.fragments),
                "continuation_confidence": table.continuation_confidence,
            }
        )
    return (
        {
            "doc_id": doc_id,
            "page": page,
            "sheet": sheet,
            "logical_tables": data_tables,
            "extraction_method": ExtractionMethod.NATIVE_TABLE
            if not sheet
            else ExtractionMethod.SPREADSHEET,
        },
        prov,
    )
