"""Spreadsheet backend (openpyxl). Values + merged ranges + formula awareness."""

from __future__ import annotations

from typing import Any, Dict, List

from .base import ParsedDocument, SheetData
from ..workspace.corpus import DocumentRef

_FORMULA_CELL_BUDGET = 20000


class SpreadsheetParser:
    name = "spreadsheet"

    def parse(self, ref: DocumentRef) -> ParsedDocument:
        try:
            import openpyxl
        except ImportError:
            return ParsedDocument(
                ref.doc_id,
                ref.fingerprint,
                ref.file_type,
                errors=["openpyxl is not installed"],
            )
        stream = ref.size_bytes > (20 * 1024 * 1024)
        try:
            wb = openpyxl.load_workbook(ref.path, read_only=stream, data_only=True)
        except Exception as exc:
            return ParsedDocument(
                ref.doc_id,
                ref.fingerprint,
                ref.file_type,
                errors=[f"workbook unreadable: {type(exc).__name__}: {exc}"],
            )
        sheets: List[SheetData] = []
        warnings: List[str] = []
        names = list(wb.sheetnames)
        for name in names:
            ws = wb[name]
            values = self._read_values(ws)
            merged = self._merged_ranges(ws)
            sheet = SheetData(
                name=name,
                max_row=ws.max_row or 0,
                max_col=ws.max_column or 0,
                values=values,
                merged=merged,
            )
            if sheet.max_row <= 1 and not values:
                warnings.append(f"sheet '{name}' is empty")
            sheets.append(sheet)
        wb.close()
        has_formulas = self._detect_formulas(ref.path, sheets)
        for sheet in sheets:
            sheet.has_formulas = has_formulas.get(sheet.name, False)
        meta = {"paginated": False, "sheets": [s.name for s in sheets]}
        return ParsedDocument(
            ref.doc_id,
            ref.fingerprint,
            ref.file_type,
            sheets=sheets,
            meta=meta,
            warnings=warnings,
        )

    @staticmethod
    def _merged_ranges(ws: Any) -> List[str]:
        """openpyxl 3.1: ws.merged_cells is a MultiCellRange (iterable)."""
        try:
            merged = getattr(ws, "merged_cells", None)
            ranges = getattr(merged, "ranges", merged)
            return [str(r) for r in ranges or []]
        except Exception:
            return []

    def _read_values(self, ws: Any) -> Dict[tuple, Any]:
        values: Dict[tuple, Any] = {}
        try:
            for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                for c_idx, value in enumerate(row, start=1):
                    if value is not None and value != "":
                        values[(r_idx, c_idx)] = value
        except Exception:
            pass
        return values

    def _detect_formulas(self, path: str, sheets: List[SheetData]) -> Dict[str, bool]:
        total = sum(s.max_row * s.max_col for s in sheets)
        if total > _FORMULA_CELL_BUDGET:
            return {}
        try:
            import openpyxl

            wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        except Exception:
            return {}
        out: Dict[str, bool] = {}
        try:
            for name in wb.sheetnames:
                found = False
                for row in wb[name].iter_rows():
                    for cell in row:
                        if cell.data_type == "f":
                            found = True
                            break
                    if found:
                        break
                out[name] = found
        finally:
            wb.close()
        return out
