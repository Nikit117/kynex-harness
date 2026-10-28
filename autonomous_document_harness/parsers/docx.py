"""DOCX backend (python-docx). Non-paginated: single logical page."""

from __future__ import annotations

from .base import ParsedDocument, ParsedPage, TableData, TextBlock
from ..workspace.corpus import DocumentRef


class DocxParser:
    name = "docx"

    def parse(self, ref: DocumentRef) -> ParsedDocument:
        try:
            import docx
        except ImportError:
            return ParsedDocument(
                ref.doc_id,
                ref.fingerprint,
                ref.file_type,
                errors=["python-docx is not installed"],
            )
        try:
            doc = docx.Document(ref.path)
        except Exception as exc:
            return ParsedDocument(
                ref.doc_id,
                ref.fingerprint,
                ref.file_type,
                errors=[f"DOCX unreadable: {type(exc).__name__}: {exc}"],
            )
        paras = [p.text for p in doc.paragraphs if p.text and p.text.strip()]
        headings = [
            p.text.strip()
            for p in doc.paragraphs
            if (p.style.name or "").startswith("Heading") and p.text.strip()
        ]
        tables = []
        for t in doc.tables:
            rows = [[c.text.strip() for c in row.cells] for row in t.rows]
            if rows:
                tables.append(TableData(rows=rows, page=1, source="docx"))
        text = "\n".join(paras)
        blocks = [TextBlock(page=1, text=p) for p in paras[:500]]
        meta = {"paginated": False, "paragraphs": len(paras), "tables": len(tables)}
        page = ParsedPage(
            page_no=1,
            text=text,
            blocks=blocks,
            tables=tables,
            headings=headings[:30],
            char_count=len(text),
        )
        return ParsedDocument(
            ref.doc_id, ref.fingerprint, ref.file_type, pages=[page], meta=meta
        )
